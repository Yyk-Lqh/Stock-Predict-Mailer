"""Bounded, read-only readers for untrusted local data files."""
from __future__ import annotations

import csv
import io
import json
import sqlite3
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

MAX_BYTES = 50 * 1024 * 1024
MAX_ROWS = 100_000
MAX_TABLES = 40
MAX_COLUMNS = 250


class IntakeError(ValueError):
    pass


@dataclass
class Table:
    name: str
    rows: list[dict]


def text_decode(raw: bytes) -> tuple[str, str]:
    encodings = ["utf-16"] if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "gb18030"]
    for encoding in encodings:
        try:
            text = raw.decode(encoding)
            if "\x00" not in text and all(ord(c) >= 32 or c in "\n\r\t" for c in text):
                return text, encoding
        except UnicodeDecodeError:
            pass
    raise IntakeError("无法可靠识别文本编码；请导出 UTF-8 文件")


def scalar(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bytes):
        return f"[binary {len(value)} bytes]"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def records(values) -> list[dict]:
    if not isinstance(values, list) or len(values) > MAX_ROWS or any(not isinstance(r, dict) for r in values):
        raise IntakeError("表格必须是对象记录列表，最多 100000 行")
    result = []
    for row in values:
        if len(row) > MAX_COLUMNS:
            raise IntakeError("字段数超过 250")
        normalized = {}
        for key, value in row.items():
            header = str(key).strip()
            if not header or header in normalized:
                raise IntakeError("存在空白或重复列名，请先明确表头")
            normalized[header] = scalar(value)
        result.append(normalized)
    return result


def matrix(name: str, values) -> Table:
    iterator = iter(values)
    header = next(iterator, None)
    if header is None:
        return Table(name, [])
    if len(header) > MAX_COLUMNS:
        raise IntakeError("字段数超过 250")
    names = [str(value).strip() if value is not None else "" for value in header]
    # Entirely empty trailing Excel columns are discarded; data under empty headers is rejected.
    while names and not names[-1]:
        names.pop()
    if not names or any(not n for n in names) or len(set(names)) != len(names):
        raise IntakeError("表头为空或重复；请使用一行明确列名")
    rows = []
    for value in iterator:
        if not any(v is not None and str(v).strip() for v in value):
            continue
        if len(value) > len(names) and any(v is not None and str(v).strip() for v in value[len(names):]):
            raise IntakeError("数据列数超过表头，请检查分隔符和空白表头")
        if len(rows) >= MAX_ROWS:
            raise IntakeError("单文件超过 100000 行，请分批导出")
        rows.append({key: scalar(value[i]) if i < len(value) else None for i, key in enumerate(names)})
    return Table(name, rows)


def read_tables(path: Path, original_name: str) -> tuple[dict, list[Table], str | None]:
    if path.stat().st_size > MAX_BYTES:
        raise IntakeError("文件超过 50MB，请分批导出")
    raw = path.read_bytes()
    suffix = Path(original_name).suffix.lower()
    if not raw:
        raise IntakeError("文件为空")
    if raw.startswith(b"PAR1"):
        import pyarrow.parquet as pq
        parquet = pq.ParquetFile(path)
        if parquet.metadata.num_rows > MAX_ROWS or parquet.metadata.num_columns > MAX_COLUMNS:
            raise IntakeError("Parquet 超过行数或列数限制")
        tables = [Table("data", records(parquet.read().to_pylist()))]
        info = {"format": "parquet"}
    elif raw.startswith(b"SQLite format 3\x00"):
        con = sqlite3.connect(path.resolve().as_uri() + "?mode=ro&immutable=1", uri=True)
        try:
            con.execute("PRAGMA query_only=ON")
            con.enable_load_extension(False)
            names = [r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")]
            if len(names) > MAX_TABLES:
                raise IntakeError("SQLite 表数超过 40，请按表导出")
            tables = []
            for name in names:
                quoted = '"' + name.replace('"', '""') + '"'
                cursor = con.execute(f"SELECT * FROM {quoted} LIMIT {MAX_ROWS + 1}")
                headers = [r[0] for r in cursor.description]
                tables.append(matrix(name, iter([headers] + cursor.fetchall())))
        finally:
            con.close()
        info = {"format": "sqlite", "read_only": True}
    elif raw.startswith(b"PK"):
        try:
            with zipfile.ZipFile(path) as archive:
                files = archive.infolist()
                if len(files) > 5000 or sum(f.file_size for f in files) > 200 * 1024 * 1024:
                    raise IntakeError("Excel 解压体积或成员数量超限")
                if any(f.file_size > max(f.compress_size, 1) * 1000 for f in files):
                    raise IntakeError("Excel 压缩比例异常")
                if "xl/workbook.xml" not in archive.namelist():
                    raise IntakeError("压缩包不是支持的 XLSX，请解压后逐个上传")
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) > MAX_TABLES:
                    raise IntakeError("Excel 工作表数超过 40")
                tables = [matrix(sheet.title, sheet.iter_rows(values_only=True)) for sheet in book.worksheets]
            finally:
                book.close()
        except zipfile.BadZipFile as exc:
            raise IntakeError("Excel 或 ZIP 文件损坏") from exc
        info = {"format": "xlsx", "formulas_evaluated": False}
    elif raw.startswith((b"%PDF", b"\x89PNG", b"\xff\xd8", b"\xd0\xcf\x11\xe0")):
        raise IntakeError("原件已归档；PDF、图片和旧 XLS 暂不自动解析，请转为文本或 XLSX")
    else:
        text, encoding = text_decode(raw)
        stripped = text.lstrip()
        if suffix in (".md", ".markdown"):
            return {"format": "markdown", "encoding": encoding}, [], text
        if suffix in (".jsonl", ".ndjson"):
            values = [json.loads(line) for line in text.splitlines() if line.strip()]
            tables, info = [Table("data", records(values))], {"format": "jsonl", "encoding": encoding}
        elif stripped.startswith(("[", "{")) or suffix == ".json":
            try:
                value = json.loads(text)
            except json.JSONDecodeError:
                values = [json.loads(line) for line in text.splitlines() if line.strip()]
                tables, info = [Table("data", records(values))], {"format": "jsonl", "encoding": encoding}
            else:
                if isinstance(value, dict):
                    wrappers = [k for k in ("data", "rows", "records", "results") if isinstance(value.get(k), list)]
                    value = value[wrappers[0]] if len(wrappers) == 1 else [value]
                tables, info = [Table("data", records(value))], {"format": "json", "encoding": encoding}
        else:
            try:
                try:
                    delimiter = csv.Sniffer().sniff(text[:32768], delimiters=",\t;|").delimiter
                except csv.Error:
                    # Sniffer can fail on short Chinese headers plus quoted thousands.
                    # Require consistent parsed field counts, never count raw commas.
                    candidates = []
                    from itertools import islice
                    for candidate in (",", "\t", ";", "|"):
                        try:
                            values = list(islice(csv.reader(io.StringIO(text), delimiter=candidate, strict=True), 50))
                        except csv.Error:
                            continue
                        widths = {len(row) for row in values if row}
                        if len(widths) == 1 and next(iter(widths)) > 1:
                            candidates.append((next(iter(widths)), candidate))
                    if not candidates:
                        raise csv.Error("no consistent delimiter")
                    candidates.sort(reverse=True)
                    if len(candidates) > 1 and candidates[0][0] == candidates[1][0]:
                        raise csv.Error("ambiguous delimiter")
                    delimiter = candidates[0][1]
                tables = [matrix("data", csv.reader(io.StringIO(text), delimiter=delimiter, strict=True))]
                info = {"format": "tsv" if delimiter == "\t" else "csv", "encoding": encoding, "delimiter": delimiter}
            except csv.Error:
                if suffix in (".csv", ".tsv"):
                    raise IntakeError("无法识别表格分隔符，请导出标准 CSV/TSV")
                return {"format": "text", "encoding": encoding}, [], text
    if sum(len(t.rows) for t in tables) > MAX_ROWS:
        raise IntakeError("所有表合计超过 100000 行，请分批导出")
    if not any(t.rows for t in tables):
        raise IntakeError("没有可处理的数据记录")
    return info, tables, None
